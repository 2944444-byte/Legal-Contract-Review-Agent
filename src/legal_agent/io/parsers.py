"""Deterministic parsers: PDF / DOCX / TXT -> normalized text + page spans.

Parsing is deterministic code (SPEC §2). We normalize whitespace lightly while
preserving character offsets closely enough that findings can cite exact
locations, and we record the character offset at which each page begins so a
:class:`~legal_agent.domain.schemas.Document` can map an offset to a page.

PDF/DOCX dependencies are imported lazily so the TXT path (and the tests that
use it) work with zero optional dependencies installed.
"""

from __future__ import annotations

import os

from ..domain.schemas import Document


class UnsupportedFormatError(ValueError):
    pass


def _normalize(text: str) -> str:
    # Normalize line endings; keep the rest intact so citations stay faithful.
    return text.replace("\r\n", "\n").replace("\r", "\n")


def _load_txt(path: str) -> Document:
    with open(path, "r", encoding="utf-8", errors="replace") as fh:
        raw = fh.read()
    text = _normalize(raw)
    return Document(
        path=path,
        text=text,
        source_format="txt",
        page_count=1,
        page_starts=[0],
        title=_guess_title(text),
    )


def _load_pdf(path: str) -> Document:
    try:
        import pdfplumber  # type: ignore
    except ImportError as exc:  # pragma: no cover - depends on optional dep
        raise UnsupportedFormatError(
            "Reading PDF requires 'pdfplumber'. Install with: pip install pdfplumber"
        ) from exc

    parts: list[str] = []
    page_starts: list[int] = []
    offset = 0
    with pdfplumber.open(path) as pdf:  # pragma: no cover - needs a real PDF
        for page in pdf.pages:
            page_starts.append(offset)
            page_text = _normalize(page.extract_text() or "")
            parts.append(page_text)
            offset += len(page_text) + 1  # +1 for the joining newline
    text = "\n".join(parts)
    return Document(
        path=path,
        text=text,
        source_format="pdf",
        page_count=max(1, len(page_starts)),
        page_starts=page_starts or [0],
        title=_guess_title(text),
    )


def _load_docx(path: str) -> Document:
    try:
        import docx  # type: ignore  # python-docx
    except ImportError as exc:  # pragma: no cover - depends on optional dep
        raise UnsupportedFormatError(
            "Reading DOCX requires 'python-docx'. Install with: pip install python-docx"
        ) from exc

    document = docx.Document(path)  # pragma: no cover - needs a real DOCX
    text = _normalize("\n".join(p.text for p in document.paragraphs))
    return Document(
        path=path,
        text=text,
        source_format="docx",
        page_count=1,  # DOCX has no fixed pagination
        page_starts=[0],
        title=_guess_title(text),
    )


def _guess_title(text: str) -> str | None:
    for line in text.splitlines():
        stripped = line.strip()
        if stripped:
            # A title is usually a short-ish first non-empty line.
            return stripped[:120]
    return None


_LOADERS = {
    ".txt": _load_txt,
    ".text": _load_txt,
    ".md": _load_txt,
    ".pdf": _load_pdf,
    ".docx": _load_docx,
}


def parse_document(path: str) -> Document:
    """Parse a contract file into a normalized :class:`Document`.

    Dispatches on file extension. Raises :class:`UnsupportedFormatError` for
    unknown extensions or missing optional dependencies.
    """
    if not os.path.exists(path):
        raise FileNotFoundError(path)
    _, ext = os.path.splitext(path.lower())
    loader = _LOADERS.get(ext)
    if loader is None:
        raise UnsupportedFormatError(
            f"Unsupported file type '{ext}'. Supported: {', '.join(sorted(_LOADERS))}"
        )
    return loader(path)
